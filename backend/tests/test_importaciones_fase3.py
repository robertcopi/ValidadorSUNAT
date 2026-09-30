import io
from decimal import Decimal
from pathlib import Path
import pytest
from httpx import AsyncClient
import openpyxl
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.consulta_cpe import ConsultaCPE
from app.services.excel_parser_service import ExcelParserService, ExcelParserException
from app.services.sunat_service import sunat_service

FIXTURES_DIR = Path(__file__).parent / "fixtures"
REAL_EXCEL_PATH = FIXTURES_DIR / "reg_compras_real.xlsx"
PRU_EXCEL_PATH = FIXTURES_DIR / "reg_compras_pru.xlsx"


async def get_auth_token(client: AsyncClient, email: str = "contador.daira@test.local") -> str:
    res = await client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "Password123!"}
    )
    assert res.status_code == 200, f"Login falló: {res.text}"
    return res.json()["access_token"]


def generar_excel_en_memoria(headers, rows, start_row=1):
    """Genera un archivo Excel (.xlsx) en memoria a partir de una lista de encabezados y filas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RegistroCompras"

    # Escribir encabezados
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=start_row, column=c_idx, value=h)

    # Escribir filas
    for r_offset, row in enumerate(rows, start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=start_row + r_offset, column=c_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# TEST 1: POST /api/v1/importaciones/preview sin JWT -> 401
# ==============================================================================
@pytest.mark.asyncio
async def test_preview_sin_jwt_rechazado(client: AsyncClient):
    files = {"archivo": ("compras.xlsx", b"dummy content", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post("/api/v1/importaciones/preview", files=files)
    assert response.status_code == 401
    assert "No se proporcionaron credenciales" in response.json()["detail"]


# ==============================================================================
# TEST 2: Archivo no .xlsx -> 400
# ==============================================================================
@pytest.mark.asyncio
async def test_preview_archivo_no_xlsx_rechazado(client: AsyncClient):
    token = await get_auth_token(client)
    files = {"archivo": ("compras.csv", b"dummy,csv,content", "text/csv")}
    response = await client.post(
        "/api/v1/importaciones/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files
    )
    assert response.status_code == 400
    assert ".xlsx" in response.json()["detail"]


# ==============================================================================
# TEST 3: Archivo vacío -> 400
# ==============================================================================
@pytest.mark.asyncio
async def test_preview_archivo_vacio_rechazado(client: AsyncClient):
    token = await get_auth_token(client)
    files = {"archivo": ("vacio.xlsx", b"", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post(
        "/api/v1/importaciones/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files
    )
    assert response.status_code == 400
    assert "vacío" in response.json()["detail"]


# ==============================================================================
# TEST 4: PRECISIÓN 9 - PREVIEW NO REALIZA CONSULTAS A SUNAT
# ==============================================================================
@pytest.mark.asyncio
async def test_preview_no_realiza_consultas_sunat(client: AsyncClient, db_session: AsyncSession, monkeypatch):
    """
    Verifica de forma estricta que POST /api/v1/importaciones/preview:
    - NO invoque sunat_service.validar_comprobante
    - NO obtenga token OAuth SUNAT
    - NO guarde registros en consultas_cpe
    """
    token = await get_auth_token(client)

    sunat_service_called = False

    async def mock_validar(*args, **kwargs):
        nonlocal sunat_service_called
        sunat_service_called = True
        raise AssertionError("¡SunatService no debe ser invocado durante el preview de Fase 3!")

    monkeypatch.setattr(sunat_service, "validar_comprobante", mock_validar)

    with open(REAL_EXCEL_PATH, "rb") as f:
        file_bytes = f.read()

    files = {"archivo": ("Reg. Compras 09-2026.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post(
        "/api/v1/importaciones/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files
    )

    assert response.status_code == 200
    assert not sunat_service_called, "SunatService fue llamado indebidamente durante el preview"

    # Verificar que no se hayan creado filas en consultas_cpe
    count_query = select(func.count(ConsultaCPE.id))
    total_consultas = (await db_session.execute(count_query)).scalar()
    assert total_consultas == 0, "No deben crearse registros en consultas_cpe durante la Fase 3"


# ==============================================================================
# TEST 5: ARCHIVO REAL 1 - Reg. Compras 09-2026.xlsx
# ==============================================================================
@pytest.mark.asyncio
async def test_archivo_real_completo(client: AsyncClient):
    """
    Prueba el archivo real Reg. Compras 09-2026.xlsx:
    - Detecta dinámicamente las 220 filas
    - Todas en estado LISTO
    - Diagnóstico de encabezados completo
    - Preserva la Nota de Crédito con monto_original negativo y monto positivo
    """
    token = await get_auth_token(client)

    with open(REAL_EXCEL_PATH, "rb") as f:
        file_bytes = f.read()

    files = {"archivo": ("Reg. Compras 09-2026.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post(
        "/api/v1/importaciones/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files
    )

    assert response.status_code == 200
    data = response.json()

    assert data["nombre_archivo"] == "Reg. Compras 09-2026.xlsx"
    assert data["total_filas_detectadas"] == 220
    assert data["total_listos"] == 220
    assert data["total_errores"] == 0
    assert data["total_duplicados"] == 0
    assert data["total_no_soportados"] == 0

    # Diagnóstico
    diag = data["diagnostico"]
    assert diag["hoja_detectada"] == "Hoja1"
    assert diag["columnas_mapeadas"]["monto"] == "Columna 22"
    assert diag["columnas_mapeadas"]["fecha_emision"] == "Columna 2"

    # Precisión 3: Verificar Nota de Crédito (fila 195)
    nc = next((f for f in data["filas"] if f["cod_comp"] == "07"), None)
    assert nc is not None, "Debe existir al menos una nota de crédito"
    assert nc["fila_excel"] == 195
    assert Decimal(str(nc["monto_original"])) == Decimal("-21142.71")
    assert Decimal(str(nc["monto"])) == Decimal("21142.71")
    assert nc["numero_serie"] == "FC01"
    assert nc["numero"] == "3242"
    assert nc["estado_archivo"] == "LISTO"
    assert nc["es_seleccionable"] is True


# ==============================================================================
# TEST 6: ARCHIVO REAL 2 - Reg. Compras 09-2026Pru.xlsx
# ==============================================================================
@pytest.mark.asyncio
async def test_archivo_real_condensado_pru(client: AsyncClient):
    """
    Prueba el archivo real Reg. Compras 09-2026Pru.xlsx:
    - Columna IMPORTE TOTAL está en la columna 11 (a diferencia de col 22)
    - El parser detecta dinámicamente la posición sin fallar
    - Mismos 220 comprobantes listos
    """
    token = await get_auth_token(client)

    with open(PRU_EXCEL_PATH, "rb") as f:
        file_bytes = f.read()

    files = {"archivo": ("Reg. Compras 09-2026Pru.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post(
        "/api/v1/importaciones/preview",
        headers={"Authorization": f"Bearer {token}"},
        files=files
    )

    assert response.status_code == 200
    data = response.json()

    assert data["total_filas_detectadas"] == 220
    assert data["total_listos"] == 220
    assert data["total_errores"] == 0
    assert data["diagnostico"]["columnas_mapeadas"]["monto"] == "Columna 11"


# ==============================================================================
# TEST 7: ENCABEZADOS DESPLAZADOS DE POSICIÓN (PRECISIÓN 6 y 12)
# ==============================================================================
def test_encabezados_desplazados():
    """Prueba que los encabezados comenzando en fila 16 sean detectados correctamente."""
    headers = [
        "CORRELATIVO",
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "PROVEEDOR RAZON SOCIAL",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["1", "15/09/2026", "01", "F001", "00001234", "6", "20538976815", "EMPRESA PRUEBA SAC", "500.00"]
    ]
    content = generar_excel_en_memoria(headers, rows, start_row=16)

    res = ExcelParserService.parse_excel(content, "desplazado.xlsx")
    assert res.total_filas_detectadas == 1
    assert res.total_listos == 1
    assert res.diagnostico.fila_inicio_encabezado == 16
    assert res.filas[0].num_ruc == "20538976815"
    assert res.filas[0].monto == Decimal("500.00")


# ==============================================================================
# TEST 8: NÚMERO Y SERIE CON CEROS A LA IZQUIERDA (PRECISIÓN 5 y 12)
# ==============================================================================
def test_preservacion_ceros_a_la_izquierda():
    """Valores como '00053117' no deben convertirse en '53117'."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["10/09/2026", "01", "0001", "00053117", "6", "20100070970", "1250.80"]
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "ceros.xlsx")
    assert res.total_filas_detectadas == 1
    fila = res.filas[0]
    assert fila.numero == "00053117", "Los ceros a la izquierda del número deben conservarse"
    assert fila.numero_serie == "0001", "Los ceros de la serie deben conservarse"
    assert fila.estado_archivo == "LISTO"


# ==============================================================================
# TEST 9: RUC ALMACENADO POR EXCEL COMO NÚMERO (PRECISIÓN 5 y 12)
# ==============================================================================
def test_ruc_almacenado_como_numero():
    """Si Excel guarda el RUC como entero 20112273922, no debe fallar ni convertirse en flotante."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    # RUC como número entero (int)
    rows = [
        ["12/09/2026", "01", "F001", "123", "6", 20112273922, 100.0]
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "ruc_num.xlsx")
    assert res.total_filas_detectadas == 1
    fila = res.filas[0]
    assert fila.num_ruc == "20112273922"
    assert fila.estado_archivo == "LISTO"


# ==============================================================================
# TEST 10: TIPO DE DOCUMENTO DE IDENTIDAD DISTINTO DE RUC (PRECISIÓN 4 y 12)
# ==============================================================================
def test_tipo_doc_identidad_distinto_de_ruc():
    """Un proveedor con DNI (tipo 1) debe quedar como NO_SOPORTADO para validación SUNAT."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["14/09/2026", "03", "B001", "45", "1", "44556677", "85.00"]
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "dni.xlsx")
    assert res.total_filas_detectadas == 1
    assert res.total_no_soportados == 1
    assert res.total_listos == 0

    fila = res.filas[0]
    assert fila.estado_archivo == "NO_SOPORTADO"
    assert fila.es_seleccionable is False
    assert any("no soportado" in err.lower() for err in fila.errores)


# ==============================================================================
# TEST 11: MONTO NEGATIVO CONSERVANDO ORIGINAL (PRECISIÓN 3 y 12)
# ==============================================================================
def test_monto_negativo_conserva_original():
    """Conserva el monto contable original negativo y normaliza en monto positivo Decimal."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["18/09/2026", "07", "FC01", "99", "6", "20538976815", -1500.50]
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "negativo.xlsx")
    fila = res.filas[0]
    assert fila.monto_original == Decimal("-1500.50")
    assert fila.monto == Decimal("1500.50")
    assert fila.estado_archivo == "LISTO"


# ==============================================================================
# TEST 12: NO CONFUNDIR CON DOCUMENTO DE REFERENCIA (PRECISIÓN 7 y 12)
# ==============================================================================
def test_distingue_comprobante_principal_de_referencia():
    """Asegura que no se tomen la serie o número del documento de referencia modificado."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
        "REFERENCIA DEL COMPROBANTE QUE SE MODIFICA TIPO",
        "REFERENCIA DEL COMPROBANTE QUE SE MODIFICA SERIE",
        "REFERENCIA DEL COMPROBANTE QUE SE MODIFICA NUMERO",
    ]
    rows = [
        ["20/09/2026", "07", "NC01", "00000001", "6", "20538976815", 350.00, "01", "F001", "00009999"]
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "referencia.xlsx")
    fila = res.filas[0]
    assert fila.cod_comp == "07"
    assert fila.numero_serie == "NC01", "Debe tomar la serie principal, no la de referencia"
    assert fila.numero == "00000001", "Debe tomar el número principal, no el de referencia"


# ==============================================================================
# TEST 13, 14, 15, 16: FILAS DE TOTALES Y REPETICIONES IGNORADAS (PRECISIÓN 12)
# ==============================================================================
def test_filas_totales_y_cabeceras_repetidas_ignoradas():
    """Omitir TOTAL PAGINA, TOTAL SUBDIARIO, TOTAL GENERAL y cabeceras intermedias."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["01/09/2026", "01", "F001", "101", "6", "20538976815", 100.00],
        ["TOTAL PAGINA", "", "", "", "", "", 100.00],
        ["FECHA DE EMISION DEL COMPROBANTE", "COMPROBANTE DE PAGO TIPO", "SERIE", "NUMERO", "", "", ""],
        ["02/09/2026", "01", "F001", "102", "6", "20538976815", 200.00],
        ["TOTAL SUBDIARIO 11", "", "", "", "", "", 300.00],
        ["TOTAL GENERAL SUBDIARIO:", "", "", "", "", "", 300.00],
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "totales.xlsx")
    assert res.total_filas_detectadas == 2, "Solo deben contarse las 2 filas de comprobantes reales"
    assert res.total_listos == 2
    assert res.filas[0].numero == "101"
    assert res.filas[1].numero == "102"


# ==============================================================================
# TEST 17: DETECCIÓN DE COMPROBANTES DUPLICADOS (PRECISIÓN 12)
# ==============================================================================
def test_deteccion_duplicados_en_archivo():
    """Si dos filas tienen el mismo RUC + Tipo + Serie + Número, la segunda es DUPLICADO."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["05/09/2026", "01", "F001", "555", "6", "20538976815", 500.00],
        ["05/09/2026", "01", "F001", "555", "6", "20538976815", 500.00],  # Idéntico
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "duplicados.xlsx")
    assert res.total_filas_detectadas == 2
    assert res.total_listos == 1
    assert res.total_duplicados == 1

    assert res.filas[0].estado_archivo == "LISTO"
    assert res.filas[0].es_seleccionable is True

    assert res.filas[1].estado_archivo == "DUPLICADO"
    assert res.filas[1].es_seleccionable is False
    assert any("duplicado" in err.lower() for err in res.filas[1].errores)


# ==============================================================================
# TEST 18: ARCHIVO MIXTO CON LISTO, ERROR, DUPLICADO Y NO_SOPORTADO
# ==============================================================================
def test_archivo_mixto_estados_completos():
    """Verifica que un archivo con variedad de registros asigne los 4 estados correctamente."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        # 1. LISTO
        ["10/09/2026", "01", "F001", "1001", "6", "20538976815", 100.00],
        # 2. DUPLICADO del anterior
        ["10/09/2026", "01", "F001", "1001", "6", "20538976815", 100.00],
        # 3. ERROR (RUC corrupto con 5 dígitos)
        ["10/09/2026", "01", "F001", "1002", "6", "12345", 200.00],
        # 4. NO_SOPORTADO (Tipo de documento 1 = DNI)
        ["10/09/2026", "01", "F001", "1003", "1", "12345678", 300.00],
        # 5. NO_SOPORTADO (Tipo de comprobante 50 = Declaración Única de Aduanas)
        ["10/09/2026", "50", "D001", "1004", "6", "20538976815", 400.00],
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "mixto.xlsx")
    assert res.total_filas_detectadas == 5
    assert res.total_listos == 1
    assert res.total_duplicados == 1
    assert res.total_errores == 1
    assert res.total_no_soportados == 2


# ==============================================================================
# TEST 19: CONVERSIÓN DIRECTA A ComprobanteValidarRequest (PRECISIÓN 10)
# ==============================================================================
def test_conversion_a_sunat_request():
    """Una fila LISTO se convierte de forma exacta al modelo ComprobanteValidarRequest de Fase 2."""
    headers = [
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "IMPORTE TOTAL",
    ]
    rows = [
        ["15/09/2026", "01", "F001", "0053117", "6", "20538976815", 1250.75]
    ]
    content = generar_excel_en_memoria(headers, rows)

    res = ExcelParserService.parse_excel(content, "req.xlsx")
    fila = res.filas[0]
    sunat_req = ExcelParserService.to_comprobante_validar_request(fila)

    assert sunat_req.num_ruc == "20538976815"
    assert sunat_req.cod_comp == "01"
    assert sunat_req.numero_serie == "F001"
    assert sunat_req.numero == "0053117"
    assert sunat_req.monto == Decimal("1250.75")
