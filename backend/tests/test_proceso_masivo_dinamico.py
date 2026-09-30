import io
from decimal import Decimal
from unittest.mock import AsyncMock, patch
import pytest
import openpyxl
from httpx import AsyncClient

from app.core.config import settings
from app.schemas.sunat import ConsultaCPEResponse, EmpresaInfo
from app.services.proceso_masivo_service import ProcesoMasivoService, proceso_masivo_service, ProcesoMasivoException


async def get_auth_token(client: AsyncClient, email: str = "contador.daira@test.local") -> str:
    """Helper para obtener token JWT de autenticación para pruebas."""
    res = await client.post(
        "/api/v1/auth/login",
        json={"email": email, "password": "Password123!"}
    )
    assert res.status_code == 200, f"Login falló: {res.text}"
    return res.json()["access_token"]


from datetime import datetime

def mock_cpe_response(num_ruc: str, cod_comp: str, serie: str, numero: str, estado: str = "VALIDO") -> ConsultaCPEResponse:
    """Crea una respuesta simulada válida de SUNAT sin llamadas reales."""
    return ConsultaCPEResponse(
        id=999,
        empresa_id=1,
        usuario_id=2,
        ruc_emisor=num_ruc,
        tipo_comprobante=cod_comp,
        tipo_comprobante_descripcion="FACTURA ELECTRÓNICA",
        serie=serie,
        numero=numero,
        fecha_emision="2026-09-15",
        monto=Decimal("150.00"),
        estado=estado,
        codigo_sunat="0",
        mensaje_sunat="Comprobante válido informado a SUNAT.",
        empresa_consultora=EmpresaInfo(ruc="20538976815", razon_social="DAIRA"),
        created_at=datetime.now(),
    )


def generar_items_comprobantes(cantidad: int):
    """Genera N items de comprobantes válidos para prueba."""
    items = []
    for i in range(1, cantidad + 1):
        num_correlativo = str(i).zfill(8)
        items.append({
            "fila_excel": i + 5,
            "num_ruc": "20100070970",
            "cod_comp": "01",
            "numero_serie": "F001",
            "numero": num_correlativo,
            "fecha_emision": "15/09/2026",
            "monto": 100.50 + i,
            "razon_social": f"PROVEEDOR TEST {i} S.A.C."
        })
    return items


def generar_excel_mes(nombre_hoja: str, cantidad_filas: int) -> bytes:
    """Genera un archivo Excel tipo Formato 8.1 con una cantidad dinámica de comprobantes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nombre_hoja

    headers = [
        "CORRELATIVO",
        "FECHA DE EMISION DEL COMPROBANTE",
        "COMPROBANTE DE PAGO TIPO",
        "COMPROBANTE DE PAGO SERIE",
        "COMPROBANTE DE PAGO NUMERO",
        "PROVEEDOR DOC. IDENTIDAD TIPO",
        "PROVEEDOR DOC. IDENTIDAD NUMERO",
        "PROVEEDOR RAZON SOCIAL",
        "IMPORTE TOTAL",
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=5, column=c_idx, value=h)

    for r_idx in range(1, cantidad_filas + 1):
        ws.cell(row=5 + r_idx, column=1, value=str(r_idx))
        ws.cell(row=5 + r_idx, column=2, value="15/09/2026")
        ws.cell(row=5 + r_idx, column=3, value="01")
        ws.cell(row=5 + r_idx, column=4, value="F001")
        ws.cell(row=5 + r_idx, column=5, value=str(r_idx).zfill(8))
        ws.cell(row=5 + r_idx, column=6, value="6")
        ws.cell(row=5 + r_idx, column=7, value="20100070970")
        ws.cell(row=5 + r_idx, column=8, value=f"PROVEEDOR {r_idx} S.A.C.")
        ws.cell(row=5 + r_idx, column=9, value=150.00 + r_idx)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# TEST 1: CÁLCULO DE PORCENTAJE DINÁMICO (CERO DEPENDENCIAS DE 220)
# ==============================================================================
def test_calculo_porcentaje_dinamico():
    """Verifica que el porcentaje de avance se calcule dinámicamente con cualquier total."""
    # 1 comprobante
    assert ProcesoMasivoService.calcular_porcentaje(1, 1) == 100.0
    # 3 comprobantes
    assert ProcesoMasivoService.calcular_porcentaje(1, 3) == 33.33
    assert ProcesoMasivoService.calcular_porcentaje(2, 3) == 66.67
    assert ProcesoMasivoService.calcular_porcentaje(3, 3) == 100.0
    # 187 comprobantes (ejemplo de la especificación)
    assert ProcesoMasivoService.calcular_porcentaje(54, 187) == 28.88
    assert ProcesoMasivoService.calcular_porcentaje(187, 187) == 100.0
    # 436 comprobantes (ejemplo de la especificación)
    assert ProcesoMasivoService.calcular_porcentaje(218, 436) == 50.0
    assert ProcesoMasivoService.calcular_porcentaje(436, 436) == 100.0
    # Casos borde
    assert ProcesoMasivoService.calcular_porcentaje(0, 100) == 0.0
    assert ProcesoMasivoService.calcular_porcentaje(0, 0) == 0.0


# ==============================================================================
# TEST 2: VALIDACIÓN CON 1 COMPROBANTE
# ==============================================================================
@pytest.mark.asyncio
async def test_proceso_masivo_1_comprobante(client: AsyncClient):
    """Verifica el flujo con exactamente 1 comprobante seleccionado."""
    token = await get_auth_token(client)
    items = generar_items_comprobantes(1)

    with patch("app.services.sunat_service.sunat_service.validar_comprobante", new_callable=AsyncMock) as mock_val:
        mock_val.return_value = mock_cpe_response("20100070970", "01", "F001", "00000001", "VALIDO")

        response = await client.post(
            "/api/v1/importaciones/validar-masivo",
            headers={"Authorization": f"Bearer {token}"},
            json={"nombre_archivo": "Registro_Compras_1_item.xlsx", "items": items}
        )

        assert response.status_code == 201
        data = response.json()
        assert data["total_registros"] == 1
        assert data["estado"] in ("PENDIENTE", "PROCESANDO", "COMPLETADO")
        proceso_id = data["id"]

        await proceso_masivo_service.esperar_proceso(proceso_id)

        res_proc = await client.get(
            f"/api/v1/procesos-masivos/{proceso_id}",
            headers={"Authorization": f"Bearer {token}"}
        )
        assert res_proc.status_code == 200
        pdata = res_proc.json()
        assert pdata["total_registros"] == 1
        assert pdata["total_procesados"] == 1
        assert pdata["porcentaje"] == 100.0
        assert pdata["total_validos"] == 1
        assert pdata["total_errores"] == 0
        assert pdata["estado"] == "COMPLETADO"


# ==============================================================================
# TEST 3: VALIDACIÓN CON 3 COMPROBANTES (CON ESTADOS DIFERENTES SEGÚN FASE 2)
# ==============================================================================
@pytest.mark.asyncio
async def test_proceso_masivo_3_comprobantes(client: AsyncClient):
    """Verifica el flujo dinámico con 3 comprobantes (VALIDO, NO_VALIDO y NO_VALIDO/ANULADO)."""
    token = await get_auth_token(client)
    items = generar_items_comprobantes(3)

    respuestas_mock = [
        mock_cpe_response("20100070970", "01", "F001", "00000001", "VALIDO"),
        mock_cpe_response("20100070970", "01", "F001", "00000002", "NO_VALIDO"),
        mock_cpe_response("20100070970", "01", "F001", "00000003", "NO_VALIDO"),
    ]

    with patch("app.services.sunat_service.sunat_service.validar_comprobante", side_effect=respuestas_mock):
        response = await client.post(
            "/api/v1/importaciones/validar-masivo",
            headers={"Authorization": f"Bearer {token}"},
            json={"nombre_archivo": "Registro_Compras_3_items.xlsx", "items": items}
        )

        assert response.status_code == 201
        proceso_id = response.json()["id"]

        await proceso_masivo_service.esperar_proceso(proceso_id)

        res_proc = await client.get(
            f"/api/v1/procesos-masivos/{proceso_id}",
            headers={"Authorization": f"Bearer {token}"}
        )
        assert res_proc.status_code == 200
        data = res_proc.json()
        assert data["total_registros"] == 3
        assert data["total_procesados"] == 3
        assert data["porcentaje"] == 100.0
        assert data["total_validos"] == 1
        assert data["total_no_validos"] == 2
        assert data["total_errores"] == 0
        assert data["estado"] == "COMPLETADO"


# ==============================================================================
# TEST 4: VALIDACIÓN CON 25 COMPROBANTES
# ==============================================================================
@pytest.mark.asyncio
async def test_proceso_masivo_25_comprobantes(client: AsyncClient):
    """Verifica el flujo dinámico con 25 comprobantes."""
    token = await get_auth_token(client)
    items = generar_items_comprobantes(25)

    with patch("app.services.sunat_service.sunat_service.validar_comprobante", new_callable=AsyncMock) as mock_val:
        mock_val.return_value = mock_cpe_response("20100070970", "01", "F001", "00000001", "VALIDO")

        response = await client.post(
            "/api/v1/importaciones/validar-masivo",
            headers={"Authorization": f"Bearer {token}"},
            json={"nombre_archivo": "Registro_Compras_25_items.xlsx", "items": items}
        )

        assert response.status_code == 201
        proceso_id = response.json()["id"]

        await proceso_masivo_service.esperar_proceso(proceso_id)

        res_proc = await client.get(
            f"/api/v1/procesos-masivos/{proceso_id}",
            headers={"Authorization": f"Bearer {token}"}
        )
        assert res_proc.status_code == 200
        data = res_proc.json()
        assert data["total_registros"] == 25
        assert data["total_procesados"] == 25
        assert data["porcentaje"] == 100.0
        assert data["total_validos"] == 25
        assert data["estado"] == "COMPLETADO"


# ==============================================================================
# TEST 5: VALIDACIÓN CON 100 COMPROBANTES
# ==============================================================================
@pytest.mark.asyncio
async def test_proceso_masivo_100_comprobantes(client: AsyncClient):
    """Verifica el flujo dinámico con 100 comprobantes."""
    token = await get_auth_token(client)
    items = generar_items_comprobantes(100)

    with patch("app.services.sunat_service.sunat_service.validar_comprobante", new_callable=AsyncMock) as mock_val:
        mock_val.return_value = mock_cpe_response("20100070970", "01", "F001", "00000001", "VALIDO")

        response = await client.post(
            "/api/v1/importaciones/validar-masivo",
            headers={"Authorization": f"Bearer {token}"},
            json={"nombre_archivo": "Registro_Compras_100_items.xlsx", "items": items}
        )

        assert response.status_code == 201
        proceso_id = response.json()["id"]

        await proceso_masivo_service.esperar_proceso(proceso_id)

        res_proc = await client.get(
            f"/api/v1/procesos-masivos/{proceso_id}",
            headers={"Authorization": f"Bearer {token}"}
        )
        assert res_proc.status_code == 200
        data = res_proc.json()
        assert data["total_registros"] == 100
        assert data["total_procesados"] == 100
        assert data["porcentaje"] == 100.0
        assert data["total_validos"] == 100
        assert data["estado"] == "COMPLETADO"


# ==============================================================================
# TEST 6: LÍMITE DE SEGURIDAD EXCEDIDO -> RECHAZO CON MENSAJE CLARO
# ==============================================================================
@pytest.mark.asyncio
async def test_limite_seguridad_excedido_rechazo_con_mensaje_claro(client: AsyncClient):
    """
    Verifica que al superar MAX_MASSIVE_ITEMS (1000 por defecto):
    - NO corta silenciosamente.
    - Rechaza con status 400.
    - Mensaje exacto formateado con comas:
      'El proceso contiene 1,245 comprobantes y supera el límite configurado de 1,000.'
    """
    token = await get_auth_token(client)
    # Generar 1,245 comprobantes
    items = generar_items_comprobantes(1245)

    response = await client.post(
        "/api/v1/importaciones/validar-masivo",
        headers={"Authorization": f"Bearer {token}"},
        json={"nombre_archivo": "Registro_Gigante.xlsx", "items": items}
    )

    assert response.status_code == 400
    data = response.json()
    assert data["detail"] == "El proceso contiene 1,245 comprobantes y supera el límite configurado de 1,000."


# ==============================================================================
# TEST 7: LÍMITE DE SEGURIDAD CONFIGURABLE MEDIANTE VARIABLE / SETTING
# ==============================================================================
@pytest.mark.asyncio
async def test_limite_seguridad_configurable_dinamicamente(client: AsyncClient):
    """
    Verifica que el límite técnico de seguridad puede modificarse (por ej. a 50)
    sin cambiar el código fuente.
    """
    token = await get_auth_token(client)
    items_55 = generar_items_comprobantes(55)

    # Modificar dinámicamente el límite a 50
    original_limit = settings.MAX_MASSIVE_ITEMS
    try:
        settings.MAX_MASSIVE_ITEMS = 50

        response = await client.post(
            "/api/v1/importaciones/validar-masivo",
            headers={"Authorization": f"Bearer {token}"},
            json={"nombre_archivo": "Test_Limite.xlsx", "items": items_55}
        )

        assert response.status_code == 400
        data = response.json()
        assert data["detail"] == "El proceso contiene 55 comprobantes y supera el límite configurado de 50."

        # Con 50 o menos debe procesar
        items_50 = generar_items_comprobantes(50)
        with patch("app.services.sunat_service.sunat_service.validar_comprobante", new_callable=AsyncMock) as mock_val:
            mock_val.return_value = mock_cpe_response("20100070970", "01", "F001", "00000001", "VALIDO")
            res_ok = await client.post(
                "/api/v1/importaciones/validar-masivo",
                headers={"Authorization": f"Bearer {token}"},
                json={"nombre_archivo": "Test_Limite_Ok.xlsx", "items": items_50}
            )
            assert res_ok.status_code == 201
            assert res_ok.json()["total_registros"] == 50
            await proceso_masivo_service.esperar_proceso(res_ok.json()["id"])
    finally:
        settings.MAX_MASSIVE_ITEMS = original_limit


# ==============================================================================
# TEST 8: ARCHIVOS DE DIFERENTES MESES (ENERO, FEBRERO, DICIEMBRE Y FUTUROS AÑOS)
# ==============================================================================
@pytest.mark.asyncio
async def test_archivos_diferentes_meses_y_anios(client: AsyncClient):
    """
    Verifica que el sistema procesa archivos de cualquier mes o año:
    - Registro de Compras Enero 2026 (7 comprobantes)
    - Registro de Compras Febrero 2026 (14 comprobantes)
    - Registro de Compras Diciembre 2026 (35 comprobantes)
    - Registro de Compras Marzo 2027 (60 comprobantes)
    El total_registros siempre proviene del contenido real, sin depender del nombre del archivo.
    """
    token = await get_auth_token(client)

    casos_meses = [
        ("Registro de Compras Enero 2026.xlsx", "Enero2026", 7),
        ("Registro de Compras Febrero 2026.xlsx", "Febrero2026", 14),
        ("Registro de Compras Diciembre 2026.xlsx", "Diciembre2026", 35),
        ("Registro de Compras Marzo 2027.xlsx", "Marzo2027", 60),
    ]

    for nombre_archivo, nombre_hoja, cantidad in casos_meses:
        excel_bytes = generar_excel_mes(nombre_hoja, cantidad)

        files = {"archivo": (nombre_archivo, excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        res = await client.post(
            "/api/v1/importaciones/preview",
            headers={"Authorization": f"Bearer {token}"},
            files=files
        )

        assert res.status_code == 200, f"Falló para {nombre_archivo}: {res.text}"
        data = res.json()
        assert data["total_filas_detectadas"] == cantidad
        assert data["total_listos"] == cantidad
        assert data["nombre_archivo"] == nombre_archivo
        assert data["diagnostico"]["hoja_detectada"] == nombre_hoja


# ==============================================================================
# TEST 9: CONSULTA DE ESTADO DE PROCESO MASIVO POR ID
# ==============================================================================
@pytest.mark.asyncio
async def test_consulta_estado_proceso_masivo_por_id(client: AsyncClient):
    """Verifica que se pueda consultar un proceso por ID mediante GET."""
    token = await get_auth_token(client)
    items = generar_items_comprobantes(5)

    with patch("app.services.sunat_service.sunat_service.validar_comprobante", new_callable=AsyncMock) as mock_val:
        mock_val.return_value = mock_cpe_response("20100070970", "01", "F001", "00000001", "VALIDO")

        post_res = await client.post(
            "/api/v1/importaciones/validar-masivo",
            headers={"Authorization": f"Bearer {token}"},
            json={"nombre_archivo": "Compras_Consulta_ID.xlsx", "items": items}
        )
        assert post_res.status_code == 201
        proceso_id = post_res.json()["id"]

        await proceso_masivo_service.esperar_proceso(proceso_id)

        get_res = await client.get(
            f"/api/v1/importaciones/proceso-masivo/{proceso_id}",
            headers={"Authorization": f"Bearer {token}"}
        )
        assert get_res.status_code == 200
        proceso_data = get_res.json()
        assert proceso_data["id"] == proceso_id
        assert proceso_data["total_registros"] == 5
        assert proceso_data["total_procesados"] == 5
        assert proceso_data["porcentaje"] == 100.0
