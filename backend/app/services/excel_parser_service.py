"""
Servicio para análisis, lectura, extracción y normalización de Registros de Compras (Formato 8.1)
en formato Excel (.xlsx).

Cumple estrictamente con las directrices de Fase 3:
- Detección dinámica de encabezados jerárquicos multinivel y resolución de celdas combinadas (merged cells).
- Clasificación de estados de archivo: LISTO, ERROR, DUPLICADO, NO_SOPORTADO (sin tocar estados SUNAT).
- Conservación de montos originales contables (Decimal negativo si aplica) y normalización positiva para SUNAT.
- Normalización no destructiva de RUC, serie y número (preservando ceros a la izquierda).
- Distinción entre comprobante principal y documentos de referencia/modificados.
- Filtrado dinámico de filas no operativas (subdiarios, subtotales, encabezados de página y totales).
- Estructuración directa para futura validación masiva en Fase 4 sin duplicidad de lógica.
"""

import io
import re
import datetime
import zipfile
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Dict, List, Optional, Tuple, Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import settings
from app.schemas.sunat import ComprobanteValidarRequest, TIPO_COMPROBANTE_MAP
from app.schemas.importacion import (
    FilaComprobantePreview,
    DiagnosticoImportacion,
    ImportacionPreviewResponse,
)


class ExcelParserException(Exception):
    """Excepción para errores de formato o estructura al procesar libros Excel."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


# Palabras clave que indican una fila no operativa (totales, cabeceras de página, etc.)
KEYWORDS_NON_VOUCHER = [
    "TOTAL PAGINA",
    "TOTAL SUBDIARIO",
    "TOTAL GENERAL",
    "TOTAL IGV",
    "TOTAL VANE",
    "FORMATO 8.",
    "PERIODO:",
    "RUC:",
    "RAZON SOCIAL:",
    "RAZÓN SOCIAL:",
    "EXPRESADO EN:",
    "SUBDIARIO :",
    "NUMERO CORRELATIVO",
    "IMPONIBLE",
    "TASA IGV",
]

# Tipos de comprobante SUNAT soportados en validación electrónica actual
TIPOS_SOPORTADOS_SUNAT = set(TIPO_COMPROBANTE_MAP.keys())  # {"01", "03", "07", "08", "14"}

# Regex de validación para RUC peruano (11 dígitos, comienza con 10, 15, 16, 17 o 20)
REGEX_RUC_VALIDO = re.compile(r"^(10|15|16|17|20)\d{9}$")


class ExcelParserService:
    """Servicio de lectura e interpretación de libros de compras Excel."""

    @staticmethod
    def validar_archivo(file_bytes: bytes, filename: str) -> None:
        """
        Valida que el archivo:
        1. Tenga extensión .xlsx
        2. No exceda el tamaño máximo configurado
        3. Tenga magic bytes de archivo ZIP (PK\\x03\\x04)
        4. No sea una bomba de descompresión (Zip Bomb / Resource Exhaustion)
        5. No contenga macros ejecutables (VBA) ni rutas con Path Traversal
        """
        clean_name = filename.strip()
        if not clean_name.lower().endswith(".xlsx"):
            raise ExcelParserException(
                message=f"Formato no soportado. El archivo '{filename}' debe ser un libro de Excel (.xlsx)."
            )

        max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
        if len(file_bytes) > max_bytes:
            raise ExcelParserException(
                message=f"El archivo excede el tamaño máximo permitido de {settings.MAX_UPLOAD_SIZE_MB}MB."
            )

        if len(file_bytes) == 0:
            raise ExcelParserException(message="El archivo subido está vacío.")

        # Validación de Magic Bytes de archivo ZIP / OOXML
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise ExcelParserException(
                message="El archivo no posee la estructura de compresión válida de un libro Excel (.xlsx)."
            )

        # Inspección de integridad y protección contra Zip Bombs
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as zf:
                infolist = zf.infolist()
                if len(infolist) > settings.MAX_EXCEL_ZIP_ENTRIES:
                    raise ExcelParserException(
                        message=f"El archivo contiene demasiadas entradas internas ({len(infolist)}). Límite de seguridad: {settings.MAX_EXCEL_ZIP_ENTRIES}."
                    )

                total_uncompressed = 0
                max_uncompressed_bytes = settings.MAX_EXCEL_UNCOMPRESSED_MB * 1024 * 1024

                for info in infolist:
                    # Protección contra Path Traversal dentro del archivo comprimido
                    if ".." in info.filename or info.filename.startswith("/") or info.filename.startswith("\\"):
                        raise ExcelParserException(
                            message="El archivo contiene nombres de ruta internos inválidos o sospechosos."
                        )

                    # Detección y rechazo de macros o binarios embebidos
                    lower_filename = info.filename.lower()
                    if lower_filename.endswith((".exe", ".vbs", ".bat", ".cmd", ".dll", ".scr", ".ps1")) or "vbaproject" in lower_filename:
                        raise ExcelParserException(
                            message="El archivo contiene macros o componentes ejecutables no autorizados."
                        )

                    total_uncompressed += info.file_size
                    if total_uncompressed > max_uncompressed_bytes:
                        raise ExcelParserException(
                            message=f"El tamaño descomprimido del archivo supera el límite de seguridad permitido ({settings.MAX_EXCEL_UNCOMPRESSED_MB}MB)."
                        )

                # Verificar ratio de compresión
                if len(file_bytes) > 0:
                    ratio = total_uncompressed / len(file_bytes)
                    if ratio > settings.MAX_EXCEL_COMPRESSION_RATIO:
                        raise ExcelParserException(
                            message="El archivo presenta un ratio de compresión anómalo que representa un riesgo de denegación de servicio."
                        )

        except zipfile.BadZipFile:
            raise ExcelParserException(message="El archivo comprimido está dañado o no es un libro Excel válido.")


    @staticmethod
    def resolver_celdas_combinadas(sheet: Worksheet, max_scan_row: int = 35) -> Dict[Tuple[int, int], Any]:
        """
        Crea una cuadrícula en memoria de las filas iniciales donde cualquier celda
        dentro de un rango combinado (merged_cells) hereda el valor de la celda superior izquierda.
        """
        grid: Dict[Tuple[int, int], Any] = {}
        max_col = sheet.max_column or 50

        # Copiar valores directos
        for r in range(1, max_scan_row + 1):
            for c in range(1, max_col + 1):
                val = sheet.cell(r, c).value
                if val is not None:
                    grid[(r, c)] = val

        # Propagar valores en rangos combinados
        for rng in sheet.merged_cells.ranges:
            if rng.min_row <= max_scan_row:
                top_val = sheet.cell(rng.min_row, rng.min_col).value
                end_r = min(rng.max_row, max_scan_row)
                end_c = min(rng.max_col, max_col)
                for r in range(rng.min_row, end_r + 1):
                    for c in range(rng.min_col, end_c + 1):
                        grid[(r, c)] = top_val

        return grid

    @classmethod
    def detectar_encabezados(
        cls, sheet: Worksheet, max_scan_row: int = 35
    ) -> Tuple[int, int, Dict[str, int]]:
        """
        Detecta dinámicamente el bloque de encabezados en el rango inicial de filas,
        resolviendo celdas combinadas y mapeando columnas mediante puntuación semántica.

        Retorna:
            (fila_inicio_encabezado, fila_fin_encabezado, mapeo_columnas)
        """
        grid = cls.resolver_celdas_combinadas(sheet, max_scan_row=max_scan_row)
        max_col = sheet.max_column or 50

        best_score = 0
        best_mapping: Dict[str, int] = {}
        best_range = (0, 0)

        # Evaluar ventanas candidatas de encabezados (entre 1 y 7 filas de grosor)
        for r_start in range(1, min(max_scan_row + 1, 30)):
            # Fila de inicio debe tener algún texto
            if not any(grid.get((r_start, c)) for c in range(1, max_col + 1)):
                continue

            for r_end in range(r_start, min(r_start + 7, max_scan_row + 1)):
                # Fila de fin debe tener algún texto
                if not any(grid.get((r_end, c)) for c in range(1, max_col + 1)):
                    continue

                col_paths: Dict[int, str] = {}

                for c in range(1, max_col + 1):
                    tokens: List[str] = []
                    for r in range(r_start, r_end + 1):
                        val = grid.get((r, c))
                        if val is not None and str(val).strip():
                            txt = str(val).strip().upper().replace("\n", " ")
                            if not tokens or tokens[-1] != txt:
                                tokens.append(txt)
                    if tokens:
                        col_paths[c] = " // ".join(tokens)

                mapping: Dict[str, int] = {}
                candidates_base: List[Tuple[int, int]] = []
                candidates_igv: List[Tuple[int, int]] = []

                for c, path in col_paths.items():
                    p = path.upper()

                    # Precisión 7: Ignorar estrictamente columnas de comprobantes modificados o referencias
                    if any(x in p for x in ["REFERENCIA", "MODIFICA", "ORIGINAL", "VENCIMIENTO", "VMTO"]):
                        continue

                    # 1. Fecha de emisión
                    if "FECHA" in p and "EMISION" in p and not any(x in p for x in ["DETRACCION"]):
                        mapping["fecha_emision"] = c

                    # 2. Tipo de comprobante (Tabla 10)
                    elif ("COMPROBANTE" in p or "DOCUMENTO" in p or "TABLA 10" in p) and "TIPO" in p and not any(x in p for x in ["PROVEEDOR", "IDENTIDAD", "CAMBIO"]):
                        mapping["cod_comp"] = c

                    # 3. Serie
                    elif ("COMPROBANTE" in p or "DOCUMENTO" in p) and "SERIE" in p and "PROVEEDOR" not in p:
                        mapping["numero_serie"] = c

                    # 4. Número de comprobante
                    elif ("COMPROBANTE" in p or "DOCUMENTO" in p) and "NUMERO" in p and not any(x in p for x in ["PROVEEDOR", "IDENTIDAD", "CORRELATIVO", "OPERACION", "DETRACCION"]):
                        mapping["numero"] = c

                    # 5. Tipo de documento de identidad del proveedor (Tabla 21)
                    elif "PROVEEDOR" in p and any(x in p for x in ["TIPO", "TABLA 21", "DOC. IDENTIDAD", "DOCUMENTO"]) and "NUMERO" not in p:
                        mapping["tipo_doc_identidad"] = c

                    # 6. Número de documento del proveedor (RUC u otro)
                    elif "PROVEEDOR" in p and any(x in p for x in ["NUMERO", "NRO", "RUC"]):
                        mapping["num_ruc"] = c

                    # 7. Razón social del proveedor
                    elif "PROVEEDOR" in p and any(x in p for x in ["APELLIDOS", "RAZON SOCIAL", "DENOMINACION", "NOMBRES"]):
                        mapping["razon_social"] = c

                    # 8. Importe total original del Registro de Compras
                    elif ("IMPORTE" in p and "TOTAL" in p) or (p.endswith("TOTAL") and not any(x in p for x in ["BASE", "IGV", "ISC", "SUBDIARIO", "GENERAL", "ADQUISICIONES", "DETRACCION"])):
                        mapping["importe_total"] = c
                        mapping["monto"] = c

                    # 9. Otros tributos y cargos (identificar para aislar y NO sumar a SUNAT)
                    elif "OTROS" in p and any(x in p for x in ["TRIBUTOS", "CARGOS"]):
                        mapping["otros_tributos"] = c

                    # 10. Tipo de cambio
                    elif ("CAMBIO" in p and any(x in p for x in ["TIPO", "T.C", "TC"])) or p.strip() in ["TIPO DE CAMBIO", "TC", "T.C."]:
                        mapping["tipo_cambio"] = c

                    # 10.1 Valor de las adquisiciones no gravadas
                    elif ("VALOR" in p and "ADQUISICIONES" in p and "NO" in p and "GRAVADAS" in p) and not any(x in p for x in ["BASE", "IMPONIBLE", "IGV", "DESTINADAS", "EXPORTACI"]):
                        mapping["adquisiciones_no_gravadas"] = c

                    # 11. Base Imponible - Grupo 1: Operaciones gravadas y/o de exportación
                    if ("BASE" in p or "IMPONIBLE" in p) and not any(x in p for x in ["NO GRAV", "IGV", "ISC", "TOTAL", "DETRACCION"]):
                        score_b = 0
                        # Grupo 1 obligatorio: Adquisiciones gravadas destinadas a operaciones gravadas y/o de exportación
                        if any(x in p for x in ["DESTINADAS A OPERACIONES", "OPERACIONES GRAVADAS", "EXPORTACI"]):
                            score_b += 20
                        if "BASE" in p and "IMPONIBLE" in p:
                            score_b += 10
                        elif "BASE" in p or "IMPONIBLE" in p:
                            score_b += 5
                        candidates_base.append((score_b, c))

                    # 12. IGV - Grupo 1: Operaciones gravadas y/o de exportación
                    if any(x in p for x in ["IGV", "I.G.V", "IMPUESTO GENERAL"]) and not any(x in p for x in ["NO GRAV", "BASE", "IMPONIBLE", "ISC", "TOTAL", "TASA", "DETRACCION"]):
                        score_i = 0
                        # Grupo 1 obligatorio
                        if any(x in p for x in ["DESTINADAS A OPERACIONES", "OPERACIONES GRAVADAS", "EXPORTACI"]):
                            score_i += 20
                        if "IGV" in p or "I.G.V" in p:
                            score_i += 10
                        elif "IMPUESTO GENERAL" in p:
                            score_i += 5
                        candidates_igv.append((score_i, c))

                if candidates_base:
                    candidates_base.sort(key=lambda x: x[0], reverse=True)
                    mapping["base_imponible"] = candidates_base[0][1]

                if candidates_igv:
                    candidates_igv.sort(key=lambda x: x[0], reverse=True)
                    mapping["igv"] = candidates_igv[0][1]

                # Evaluar presencia de columnas críticas
                score = 0
                for f in ["fecha_emision", "cod_comp", "numero_serie", "numero", "num_ruc"]:
                    if f in mapping:
                        score += 2
                if "base_imponible" in mapping and "igv" in mapping:
                    score += 4
                elif "importe_total" in mapping or "monto" in mapping:
                    score += 2
                if "razon_social" in mapping:
                    score += 1
                if "tipo_doc_identidad" in mapping:
                    score += 1
                if "tipo_cambio" in mapping:
                    score += 1
                if "adquisiciones_no_gravadas" in mapping:
                    score += 1

                # Priorizar coincidencia completa con mayor detalle y menor grosor superfluo
                span = r_end - r_start
                best_span = best_range[1] - best_range[0] if best_range != (0, 0) else 999
                if (score > best_score) or \
                   (score == best_score and len(mapping) > len(best_mapping)) or \
                   (score == best_score and len(mapping) == len(best_mapping) and span < best_span):
                    best_score = score
                    best_mapping = mapping
                    best_range = (r_start, r_end)

        if best_score < 4:
            raise ExcelParserException(
                message="No se pudo identificar la estructura del Registro de Compras en el archivo. Verifique que corresponda al Formato 8.1."
            )

        return best_range[0], best_range[1], best_mapping

    @classmethod
    def es_fila_no_comprobante(
        cls,
        sheet: Worksheet,
        r: int,
        max_col: int,
        mapping: Dict[str, int],
    ) -> bool:
        """
        Determina si una fila corresponde a un pie de página, subtotal, subdiario,
        repetición de cabecera o fila vacía que debe ser omitida del análisis de comprobantes.
        """
        vals = [sheet.cell(r, c).value for c in range(1, max_col + 1)]
        if not any(v is not None and str(v).strip() for v in vals):
            return True

        row_text = " ".join(str(v or "").strip() for v in vals).upper()

        # Filtrar palabras clave no operativas
        if any(kw in row_text for kw in KEYWORDS_NON_VOUCHER):
            return True

        # Filtrar repeticiones de cabecera en saltos de página
        c1 = str(sheet.cell(r, 1).value or "").strip().upper()
        if any(h in c1 for h in ["NUMERO", "CORRELATIVO", "CODIGO UNICO", "DEL REGISTRO", "DE LA OPERACION"]):
            return True

        fecha_col = mapping.get("fecha_emision")
        if fecha_col:
            c_fecha = str(sheet.cell(r, fecha_col).value or "").strip().upper()
            if any(h in c_fecha for h in ["FECHA", "EMISION", "COMPROBANTE", "DE PAGO", "VMTO"]):
                return True

        tipo_col = mapping.get("cod_comp")
        if tipo_col:
            c_tipo = str(sheet.cell(r, tipo_col).value or "").strip().upper()
            if any(h in c_tipo for h in ["COMPROBANTE", "DOCUMENTO", "TIPO"]):
                return True

        return False

    @classmethod
    def normalizar_celda_texto(cls, valor: Any) -> str:
        """
        Normaliza el contenido de una celda a cadena sin pérdida destructiva de ceros a la izquierda.
        """
        if valor is None:
            return ""
        if isinstance(valor, float):
            if valor.is_integer():
                return str(int(valor))
            return str(valor).strip()
        if isinstance(valor, int):
            return str(valor)
        return str(valor).strip()

    @classmethod
    def normalizar_fecha(cls, valor: Any) -> Tuple[Optional[str], Optional[str]]:
        """
        Normaliza una fecha al formato DD/MM/YYYY requerido por SUNAT.
        Retorna: (fecha_formateada, error_si_hubo)
        """
        if valor is None:
            return None, "Fecha de emisión vacía."

        if isinstance(valor, (datetime.date, datetime.datetime)):
            return valor.strftime("%d/%m/%Y"), None

        texto = str(valor).strip()
        if not texto:
            return None, "Fecha de emisión vacía."

        # Probar formatos comunes: DD/MM/YYYY, YYYY-MM-DD, DD-MM-YYYY
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                dt = datetime.datetime.strptime(texto, fmt)
                return dt.strftime("%d/%m/%Y"), None
            except ValueError:
                continue

        return None, f"Formato de fecha no reconocido: '{texto}'. Se esperaba DD/MM/YYYY."

    @classmethod
    def parse_decimal_monetario(
        cls, valor: Any, permitir_vacio: bool = True
    ) -> Tuple[Optional[Decimal], Optional[str]]:
        """
        Convierte una celda a Decimal con redondeo a 2 decimales sin usar float.
        Si permitir_vacio es True y el valor es None o vacío, retorna (Decimal('0.00'), None).
        """
        if valor is None:
            return (Decimal("0.00"), None) if permitir_vacio else (None, "Valor monetario vacío.")
        texto = str(valor).strip().replace(",", "")
        if not texto:
            return (Decimal("0.00"), None) if permitir_vacio else (None, "Valor monetario vacío.")
        try:
            val = Decimal(texto).quantize(Decimal("0.01"))
            return val, None
        except (InvalidOperation, ValueError):
            return None, f"Valor monetario numérico inválido: '{texto}'."

    @classmethod
    def normalizar_monto(cls, valor: Any) -> Tuple[Optional[Decimal], Optional[Decimal], Optional[str]]:
        """
        Normaliza el monto contable respetando la Precisión 3:
        - monto_original: conserva signo contable original (puede ser negativo ej. notas de crédito).
        - monto: valor absoluto positivo para consulta a SUNAT.
        Retorna: (monto_original, monto_normalizado, error_si_hubo)
        """
        if valor is None:
            return None, None, "Monto total vacío."

        texto = str(valor).strip().replace(",", "")
        if not texto:
            return None, None, "Monto total vacío."

        try:
            monto_original = Decimal(texto).quantize(Decimal("0.01"))
            monto_normalizado = abs(monto_original)
            return monto_original, monto_normalizado, None
        except (InvalidOperation, ValueError):
            return None, None, f"Valor de monto numérico inválido: '{texto}'."

    @classmethod
    def parse_excel(cls, file_bytes: bytes, filename: str) -> ImportacionPreviewResponse:
        """
        Punto de entrada principal para el análisis del libro Excel.
        Ejecuta: validación -> detección de estructura -> extracción -> normalización ->
        detección de duplicados -> clasificación de filas sin invocar servicios externos.
        """
        cls.validar_archivo(file_bytes, filename)

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as e:
            raise ExcelParserException(message=f"No se pudo abrir el archivo Excel: {str(e)}")

        sheet = wb.active
        if sheet is None:
            raise ExcelParserException(message="El archivo Excel no contiene hojas de cálculo activas.")

        if (sheet.max_row or 0) > settings.MAX_EXCEL_ROWS:
            raise ExcelParserException(
                message=f"El archivo contiene demasiadas filas ({sheet.max_row}). Límite de seguridad: {settings.MAX_EXCEL_ROWS}."
            )
        if (sheet.max_column or 0) > settings.MAX_EXCEL_COLS:
            raise ExcelParserException(
                message=f"El archivo contiene demasiadas columnas ({sheet.max_column}). Límite de seguridad: {settings.MAX_EXCEL_COLS}."
            )

        max_scan = min(sheet.max_row or 1, 35)
        fila_inicio, fila_fin, mapping = cls.detectar_encabezados(sheet, max_scan_row=max_scan)

        # Mapeo descriptivo para diagnóstico
        col_names_diag = {
            "fecha_emision": f"Columna {mapping.get('fecha_emision')}",
            "cod_comp": f"Columna {mapping.get('cod_comp')}",
            "numero_serie": f"Columna {mapping.get('numero_serie')}",
            "numero": f"Columna {mapping.get('numero')}",
            "tipo_doc_identidad": f"Columna {mapping.get('tipo_doc_identidad')}" if "tipo_doc_identidad" in mapping else "No detectada",
            "num_ruc": f"Columna {mapping.get('num_ruc')}",
            "razon_social": f"Columna {mapping.get('razon_social')}" if "razon_social" in mapping else "No detectada",
            "base_imponible": f"Columna {mapping.get('base_imponible')}" if "base_imponible" in mapping else "No detectada",
            "igv": f"Columna {mapping.get('igv')}" if "igv" in mapping else "No detectada",
            "adquisiciones_no_gravadas": f"Columna {mapping.get('adquisiciones_no_gravadas')}" if "adquisiciones_no_gravadas" in mapping else "No detectada",
            "tipo_cambio": f"Columna {mapping.get('tipo_cambio')}" if "tipo_cambio" in mapping else "No detectada",
            "importe_total": f"Columna {mapping.get('importe_total') or mapping.get('monto')}" if ("importe_total" in mapping or "monto" in mapping) else "No detectada",
            "monto": f"Columna {mapping.get('monto') or mapping.get('importe_total')}",
        }

        diagnostico = DiagnosticoImportacion(
            hoja_detectada=sheet.title,
            fila_inicio_encabezado=fila_inicio,
            fila_fin_encabezado=fila_fin,
            columnas_mapeadas=col_names_diag,
        )

        filas_preview: List[FilaComprobantePreview] = []
        # Registro para detección de duplicados internos en el archivo:
        # Clave: (num_ruc, cod_comp, numero_serie, numero) -> primera fila_excel observada
        vistos_en_archivo: Dict[Tuple[str, str, str, str], int] = {}

        total_listos = 0
        total_errores = 0
        total_duplicados = 0
        total_no_soportados = 0

        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 50

        # Iterar filas a partir del final del bloque de encabezados
        for r in range(fila_fin + 1, max_row + 1):
            if cls.es_fila_no_comprobante(sheet, r, max_col, mapping):
                continue

            errores_fila: List[str] = []
            es_no_soportado = False

            # Extracción cruda respetando ceros iniciales
            raw_fecha = sheet.cell(r, mapping["fecha_emision"]).value if "fecha_emision" in mapping else None
            raw_tipo = sheet.cell(r, mapping["cod_comp"]).value if "cod_comp" in mapping else None
            raw_serie = sheet.cell(r, mapping["numero_serie"]).value if "numero_serie" in mapping else None
            raw_num = sheet.cell(r, mapping["numero"]).value if "numero" in mapping else None
            raw_tipo_doc = sheet.cell(r, mapping["tipo_doc_identidad"]).value if "tipo_doc_identidad" in mapping else None
            raw_ruc = sheet.cell(r, mapping["num_ruc"]).value if "num_ruc" in mapping else None
            raw_razon = sheet.cell(r, mapping["razon_social"]).value if "razon_social" in mapping else None
            raw_base = sheet.cell(r, mapping["base_imponible"]).value if "base_imponible" in mapping else None
            raw_igv = sheet.cell(r, mapping["igv"]).value if "igv" in mapping else None
            raw_no_gravadas = sheet.cell(r, mapping["adquisiciones_no_gravadas"]).value if "adquisiciones_no_gravadas" in mapping else None
            raw_tc = sheet.cell(r, mapping["tipo_cambio"]).value if "tipo_cambio" in mapping else None
            raw_total = sheet.cell(r, mapping.get("importe_total") or mapping.get("monto")).value if ("importe_total" in mapping or "monto" in mapping) else None
            raw_otros = sheet.cell(r, mapping["otros_tributos"]).value if "otros_tributos" in mapping else None

            # 1. Normalización de Tipo de Documento de Identidad (Precisión 4)
            tipo_doc_clean = cls.normalizar_celda_texto(raw_tipo_doc)
            num_ruc_clean = cls.normalizar_celda_texto(raw_ruc).replace("-", "").replace(" ", "")

            if tipo_doc_clean and tipo_doc_clean != "6":
                es_no_soportado = True
                errores_fila.append(
                    f"Tipo de documento de identidad '{tipo_doc_clean}' no soportado (solo se valida RUC tipo 6 en SUNAT)."
                )
            elif not num_ruc_clean:
                errores_fila.append("Número de RUC del proveedor no encontrado en la fila.")
            elif not REGEX_RUC_VALIDO.match(num_ruc_clean):
                errores_fila.append(
                    f"RUC inválido: '{num_ruc_clean}'. Debe tener 11 dígitos numéricos y comenzar con 10, 15, 16, 17 o 20."
                )

            # 2. Normalización de Tipo de Comprobante
            tipo_clean = cls.normalizar_celda_texto(raw_tipo)
            if tipo_clean:
                cod_comp = tipo_clean.zfill(2)
            else:
                cod_comp = ""
                errores_fila.append("Tipo de comprobante vacío.")

            if cod_comp:
                if cod_comp in TIPOS_SOPORTADOS_SUNAT:
                    tipo_desc = TIPO_COMPROBANTE_MAP[cod_comp]
                else:
                    es_no_soportado = True
                    tipo_desc = f"Tipo {cod_comp} (No soportado)"
                    errores_fila.append(
                        f"Tipo de comprobante '{cod_comp}' no soportado para validación electrónica SUNAT."
                    )
            else:
                tipo_desc = "Desconocido"

            # 3. Normalización de Serie y Número (Precisión 5: conservación de ceros a la izquierda)
            serie_clean = cls.normalizar_celda_texto(raw_serie)
            num_clean = cls.normalizar_celda_texto(raw_num)

            if not serie_clean:
                errores_fila.append("Serie del comprobante vacía.")
            if not num_clean:
                errores_fila.append("Número correlativo del comprobante vacío.")

            # 4. Normalización de Fecha
            fecha_norm, err_fecha = cls.normalizar_fecha(raw_fecha)
            if err_fecha:
                errores_fila.append(err_fecha)

            # 5. Normalización y Cálculo de Montos (MONTO SUNAT = (BASE IMPONIBLE + IGV + NO GRAVADAS) [/ TIPO CAMBIO])
            base_imponible_dec: Optional[Decimal] = None
            igv_dec: Optional[Decimal] = None
            monto_base_igv: Optional[Decimal] = None
            valor_adquisiciones_no_gravadas: Optional[Decimal] = None
            monto_calculado_original: Optional[Decimal] = None
            tipo_cambio_dec: Optional[Decimal] = None
            monto_convertido_dec: Optional[Decimal] = None
            importe_total_excel: Optional[Decimal] = None
            otros_tributos_dec: Optional[Decimal] = None
            monto_orig: Optional[Decimal] = None
            monto_norm: Optional[Decimal] = None

            # Extracción y validación segura del Tipo de Cambio
            if raw_tc is not None and str(raw_tc).strip():
                tc_str = str(raw_tc).strip().replace(",", "")
                try:
                    val_tc = Decimal(tc_str)
                    if val_tc > Decimal("0.00"):
                        tipo_cambio_dec = val_tc
                    elif val_tc == Decimal("0.00"):
                        tipo_cambio_dec = Decimal("0.00")
                except (InvalidOperation, ValueError):
                    errores_fila.append(f"Tipo de cambio inválido o no numérico: '{raw_tc}'.")

            if "base_imponible" in mapping and "igv" in mapping:
                b_str = str(raw_base).strip() if raw_base is not None else ""
                i_str = str(raw_igv).strip() if raw_igv is not None else ""
                t_str = str(raw_total).strip() if raw_total is not None else ""

                if not b_str and not i_str and not t_str:
                    errores_fila.append("Montos vacíos: Base Imponible e IGV no encontrados.")
                else:
                    base_imponible_dec, err_b = cls.parse_decimal_monetario(raw_base, permitir_vacio=True)
                    igv_dec, err_i = cls.parse_decimal_monetario(raw_igv, permitir_vacio=True)

                    if err_b:
                        errores_fila.append(f"Base Imponible inválida: {err_b}")
                    if err_i:
                        errores_fila.append(f"IGV inválido: {err_i}")

                    # Extraer Valor de las Adquisiciones No Gravadas de la MISMA FILA
                    no_gravadas_calc = Decimal("0.00")
                    err_ng = None
                    if raw_no_gravadas is not None and str(raw_no_gravadas).strip():
                        val_ng, err_ng = cls.parse_decimal_monetario(raw_no_gravadas, permitir_vacio=False)
                        if err_ng:
                            errores_fila.append(f"Valor de Adquisiciones No Gravadas inválido: {err_ng}")
                        else:
                            valor_adquisiciones_no_gravadas = val_ng
                            no_gravadas_calc = val_ng
                    else:
                        valor_adquisiciones_no_gravadas = None
                        no_gravadas_calc = Decimal("0.00")

                    # Extraer Importe Total del Excel si existe la columna (para visualización e historial)
                    if t_str:
                        importe_total_excel, _ = cls.parse_decimal_monetario(raw_total, permitir_vacio=False)

                    # Extraer Otros tributos y cargos si existe la columna (NO se suma al monto SUNAT)
                    if raw_otros is not None and str(raw_otros).strip():
                        otros_tributos_dec, _ = cls.parse_decimal_monetario(raw_otros, permitir_vacio=False)

                    if base_imponible_dec is not None and igv_dec is not None and not err_b and not err_i and not err_ng:
                        # REGLA 1: monto_base_igv = BASE IMPONIBLE + IGV (Decimal estricto)
                        monto_base_igv = (base_imponible_dec + igv_dec).quantize(Decimal("0.01"))

                        # REGLA 2: monto_calculado_original = BASE + IGV + NO GRAVADAS
                        monto_calculado_original = (monto_base_igv + no_gravadas_calc).quantize(Decimal("0.01"))

                        # REGLA 3: APLICAR TIPO DE CAMBIO SI EXISTE Y ES > 0
                        if tipo_cambio_dec is not None and tipo_cambio_dec > Decimal("0.00"):
                            # monto_convertido = (monto_calculado_original / tipo_cambio) con redondeo HALF_UP a 2 decimales
                            monto_convertido_dec = (monto_calculado_original / tipo_cambio_dec).quantize(
                                Decimal("0.01"), rounding=ROUND_HALF_UP
                            )
                            # Para SUNAT se envía el valor absoluto positivo
                            monto_norm = abs(monto_convertido_dec)
                            # Para trazabilidad contable original se conserva el signo (ej. notas de crédito negativas)
                            monto_orig = monto_convertido_dec
                        else:
                            # Sin tipo de cambio o tipo de cambio <= 0: no dividir
                            monto_convertido_dec = None
                            monto_norm = abs(monto_calculado_original)
                            monto_orig = monto_calculado_original
            else:
                # Fallback para archivos sin columnas multinivel de Base Imponible e IGV (ej. condensados)
                monto_orig, monto_norm, err_monto = cls.normalizar_monto(raw_total)
                if err_monto:
                    errores_fila.append(err_monto)
                else:
                    monto_base_igv = monto_orig
                    monto_calculado_original = monto_orig
                    importe_total_excel = monto_orig
                    if tipo_cambio_dec is not None and tipo_cambio_dec > Decimal("0.00"):
                        monto_convertido_dec = (monto_calculado_original / tipo_cambio_dec).quantize(
                            Decimal("0.01"), rounding=ROUND_HALF_UP
                        )
                        monto_norm = abs(monto_convertido_dec)
                        monto_orig = monto_convertido_dec

            razon_social_clean = cls.normalizar_celda_texto(raw_razon) or None

            # 6. Detección de Duplicados Internos
            es_duplicado = False
            if num_ruc_clean and cod_comp and serie_clean and num_clean and not errores_fila and not es_no_soportado:
                clave_cpe = (num_ruc_clean, cod_comp, serie_clean.upper(), num_clean)
                if clave_cpe in vistos_en_archivo:
                    es_duplicado = True
                    fila_previa = vistos_en_archivo[clave_cpe]
                    errores_fila.append(
                        f"Comprobante duplicado en el archivo (coincide con la fila {fila_previa})."
                    )
                else:
                    vistos_en_archivo[clave_cpe] = r

            # 7. Asignación de Estado de Archivo (Precisión 8: nunca estado SUNAT en Fase 3)
            if es_duplicado:
                estado_archivo = "DUPLICADO"
                total_duplicados += 1
                es_seleccionable = False
            elif es_no_soportado:
                estado_archivo = "NO_SOPORTADO"
                total_no_soportados += 1
                es_seleccionable = False
            elif errores_fila:
                estado_archivo = "ERROR"
                total_errores += 1
                es_seleccionable = False
            else:
                estado_archivo = "LISTO"
                total_listos += 1
                es_seleccionable = True

            filas_preview.append(
                FilaComprobantePreview(
                    fila_excel=r,
                    num_ruc=num_ruc_clean or None,
                    tipo_doc_identidad=tipo_doc_clean or None,
                    cod_comp=cod_comp or None,
                    tipo_descripcion=tipo_desc,
                    numero_serie=serie_clean or None,
                    numero=num_clean or None,
                    fecha_emision=fecha_norm,
                    monto_original=monto_orig,
                    monto=monto_norm,
                    razon_social_emisor=razon_social_clean,
                    estado_archivo=estado_archivo,
                    errores=errores_fila,
                    es_seleccionable=es_seleccionable,
                    base_imponible=base_imponible_dec,
                    igv=igv_dec,
                    monto_base_igv=monto_base_igv,
                    valor_adquisiciones_no_gravadas=valor_adquisiciones_no_gravadas,
                    monto_calculado_original=monto_calculado_original,
                    tipo_cambio=tipo_cambio_dec,
                    monto_convertido=monto_convertido_dec,
                    importe_total_excel=importe_total_excel,
                    otros_tributos=otros_tributos_dec,
                )
            )

        total_filas = len(filas_preview)

        return ImportacionPreviewResponse(
            nombre_archivo=filename,
            diagnostico=diagnostico,
            total_filas_detectadas=total_filas,
            total_listos=total_listos,
            total_errores=total_errores,
            total_duplicados=total_duplicados,
            total_no_soportados=total_no_soportados,
            filas=filas_preview,
        )

    @staticmethod
    def to_comprobante_validar_request(fila: FilaComprobantePreview) -> ComprobanteValidarRequest:
        """
        Convierte una fila en estado LISTO directamente al modelo de request
        validado en Fase 2 (preparación para Fase 4, Precisión 10).
        """
        if not fila.es_seleccionable or fila.estado_archivo != "LISTO":
            raise ExcelParserException(
                message=f"La fila {fila.fila_excel} no está en estado LISTO para ser consultada."
            )

        return ComprobanteValidarRequest(
            num_ruc=fila.num_ruc,  # type: ignore
            cod_comp=fila.cod_comp,  # type: ignore
            numero_serie=fila.numero_serie,  # type: ignore
            numero=fila.numero,  # type: ignore
            fecha_emision=fila.fecha_emision,  # type: ignore
            monto=fila.monto,  # type: ignore
        )
